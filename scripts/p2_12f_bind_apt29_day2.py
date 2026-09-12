from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

TECHNIQUE_RE = re.compile(r"\bT\d{4}(?:\.\d{3})?\b", re.IGNORECASE)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def event_source(doc: dict) -> dict:
    src = doc.get("_source")
    return src if isinstance(src, dict) else doc


def event_channel(src: dict) -> str | None:
    for key in ("Channel", "channel"):
        value = src.get(key)
        if value not in (None, ""):
            return str(value)
    winlog = src.get("winlog")
    if isinstance(winlog, dict) and winlog.get("channel") not in (None, ""):
        return str(winlog["channel"])
    return None


def event_id(src: dict) -> str | None:
    for key in ("EventID", "event_id"):
        value = src.get(key)
        if value not in (None, ""):
            return str(value)
    winlog = src.get("winlog")
    if isinstance(winlog, dict) and winlog.get("event_id") not in (None, ""):
        return str(winlog["event_id"])
    event = src.get("event")
    if isinstance(event, dict) and event.get("code") not in (None, ""):
        return str(event["code"])
    return None


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--archive", required=True)
    parser.add_argument("--plan", required=True)
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    archive = Path(args.archive)
    plan = Path(args.plan)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    assert archive.stat().st_size == 43_033_041, archive.stat().st_size
    assert plan.stat().st_size == 25_051, plan.stat().st_size

    with zipfile.ZipFile(archive) as zf:
        assert zf.testzip() is None
        infos = zf.infolist()
        assert len(infos) == 1, [x.filename for x in infos]
        info = infos[0]
        total_events = 0
        parse_errors = 0
        channels: Counter[str] = Counter()
        sysmon_event_ids: Counter[str] = Counter()
        with zf.open(info) as raw:
            for line in raw:
                if not line.strip():
                    continue
                try:
                    doc = json.loads(line)
                except Exception:
                    parse_errors += 1
                    continue
                total_events += 1
                src = event_source(doc)
                channel = event_channel(src)
                eid = event_id(src)
                if channel:
                    channels[channel] += 1
                if channel and "sysmon" in channel.lower() and eid:
                    sysmon_event_ids[eid] += 1

    wb = load_workbook(plan, read_only=True, data_only=True)
    day2_sheet = None
    for name in wb.sheetnames:
        if name.strip().lower() == "day2":
            day2_sheet = name
            break
    assert day2_sheet is not None, wb.sheetnames
    ws = wb[day2_sheet]

    labeled_rows = []
    techniques: dict[str, list[int]] = {}
    for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        values = [None if value is None else str(value) for value in row]
        if not any(value not in (None, "") for value in values):
            continue
        joined = " | ".join(value or "" for value in values)
        found = sorted({m.group(0).upper() for m in TECHNIQUE_RE.finditer(joined)})
        if found:
            labeled_rows.append({"row": row_idx, "values": values, "techniques": found})
            for technique in found:
                techniques.setdefault(technique, []).append(row_idx)

    result = {
        "schema": "breachscope.p2_12f_apt29_day2_binding_probe.v1",
        "evaluation_class": "confirmatory_same_campaign_holdout_pre_scoring",
        "source": {
            "repository": "OTRF/Security-Datasets",
            "pinned_commit": "d9d40ef123d2c87d5d3df28c96bcab4f0faccc87",
            "archive_path": "datasets/compound/apt29/day2/apt29_evals_day2_manual.zip",
            "archive_git_blob_sha1": "15e1e9d2d88a729b832c6430fa285c3e87bb70e4",
            "plan_path": "datasets/compound/apt29/emulationplans/apt29.xlsx",
            "plan_git_blob_sha1": "e2b95dc306967a6a2d9af033cf1fe7c3ad155c13",
        },
        "archive": {
            "size_bytes": archive.stat().st_size,
            "sha256": sha256(archive),
            "member_count": 1,
            "member_name": info.filename,
            "member_compressed_bytes": info.compress_size,
            "member_uncompressed_bytes": info.file_size,
            "member_crc32": f"{info.CRC:08x}",
            "actual_parse": {
                "total_events": total_events,
                "parse_errors": parse_errors,
                "channels": dict(sorted(channels.items())),
                "sysmon_event_ids": dict(sorted(sysmon_event_ids.items(), key=lambda x: int(x[0]) if x[0].isdigit() else 999999)),
            },
        },
        "emulation_plan": {
            "size_bytes": plan.stat().st_size,
            "sha256": sha256(plan),
            "sheet_names": wb.sheetnames,
            "selected_sheet": day2_sheet,
            "labeled_row_count": len(labeled_rows),
            "unique_technique_count": len(techniques),
            "technique_ids": sorted(techniques),
            "technique_occurrences": techniques,
            "labeled_rows": labeled_rows,
        },
        "protocol": {
            "detector_executed": False,
            "detector_imported": False,
            "detector_results_consulted": False,
            "day1_result_already_observed": True,
            "purpose": "Day 2 byte binding, actual archive parse, and source-label extraction only",
        },
        "claim_boundary": {
            "independent_fresh_external_holdout": "NOT_CLAIMED",
            "confirmatory_holdout_result": "NOT_YET_MEASURED",
            "production_detection_rate": "NOT_CLAIMED",
            "production_precision": "NOT_CLAIMED",
            "production_recall": "NOT_CLAIMED",
            "production_false_positive_rate": "NOT_CLAIMED",
        },
    }

    output = out / "binding-probe.json"
    output.write_text(json.dumps(result, indent=2, ensure_ascii=False, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps({
        "archive_sha256": result["archive"]["sha256"],
        "archive_size_bytes": result["archive"]["size_bytes"],
        "member_name": result["archive"]["member_name"],
        "member_uncompressed_bytes": result["archive"]["member_uncompressed_bytes"],
        "total_events": total_events,
        "parse_errors": parse_errors,
        "plan_sha256": result["emulation_plan"]["sha256"],
        "day2_labeled_rows": len(labeled_rows),
        "day2_unique_techniques": len(techniques),
        "technique_ids": sorted(techniques),
        "detector_executed": False,
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
