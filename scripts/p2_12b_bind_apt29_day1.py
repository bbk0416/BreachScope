from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

TECHNIQUE_RE = re.compile(r"\bT\d{4}(?:\.\d{3})?\b", re.IGNORECASE)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def main() -> int:
    p = argparse.ArgumentParser()
    p.add_argument("--archive", required=True)
    p.add_argument("--plan", required=True)
    p.add_argument("--out", required=True)
    args = p.parse_args()

    archive = Path(args.archive)
    plan = Path(args.plan)
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    assert archive.stat().st_size == 13_944_973, archive.stat().st_size
    assert plan.stat().st_size == 25_051, plan.stat().st_size

    ext_counts: Counter[str] = Counter()
    entries = []
    total_uncompressed = 0
    with zipfile.ZipFile(archive) as zf:
        bad = zf.testzip()
        assert bad is None, bad
        for info in zf.infolist():
            suffix = Path(info.filename).suffix.lower() or "<none>"
            ext_counts[suffix] += 1
            total_uncompressed += info.file_size
            entries.append({
                "name": info.filename,
                "compressed_size": info.compress_size,
                "uncompressed_size": info.file_size,
                "crc32": f"{info.CRC:08x}",
            })

    wb = load_workbook(plan, read_only=True, data_only=True)
    plan_rows = []
    techniques: dict[str, list[dict[str, object]]] = {}
    for ws in wb.worksheets:
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = [None if v is None else str(v) for v in row]
            if not any(v not in (None, "") for v in values):
                continue
            joined = " | ".join(v or "" for v in values)
            found = sorted({m.group(0).upper() for m in TECHNIQUE_RE.finditer(joined)})
            record = {"sheet": ws.title, "row": row_idx, "values": values, "techniques": found}
            plan_rows.append(record)
            for technique in found:
                techniques.setdefault(technique, []).append({"sheet": ws.title, "row": row_idx})

    binding = {
        "schema": "breachscope.p2_12b_apt29_day1_binding_probe.v1",
        "source": {
            "repository": "OTRF/Security-Datasets",
            "pinned_commit": "d9d40ef123d2c87d5d3df28c96bcab4f0faccc87",
            "archive_path": "datasets/compound/apt29/day1/apt29_evals_day1_manual.zip",
            "archive_git_blob_sha1": "7352679a173ec0310f9d0ed587782545182dd394",
            "plan_path": "datasets/compound/apt29/emulationplans/apt29.xlsx",
            "plan_git_blob_sha1": "e2b95dc306967a6a2d9af033cf1fe7c3ad155c13",
        },
        "archive": {
            "size_bytes": archive.stat().st_size,
            "sha256": sha256(archive),
            "zip_entry_count": len(entries),
            "total_uncompressed_bytes": total_uncompressed,
            "extension_counts": dict(sorted(ext_counts.items())),
            "entries": entries,
        },
        "emulation_plan": {
            "size_bytes": plan.stat().st_size,
            "sha256": sha256(plan),
            "sheet_names": wb.sheetnames,
            "nonempty_row_count": len(plan_rows),
            "technique_ids_found": sorted(techniques),
            "technique_occurrences": techniques,
        },
        "protocol": {
            "detector_executed": False,
            "detector_results_consulted": False,
            "purpose": "byte binding and source-label extraction only",
        },
    }

    (out / "binding-probe.json").write_text(
        json.dumps(binding, indent=2, ensure_ascii=False, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    (out / "emulation-plan-rows.json").write_text(
        json.dumps({"rows": plan_rows}, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    print(json.dumps({
        "archive_sha256": binding["archive"]["sha256"],
        "archive_size_bytes": binding["archive"]["size_bytes"],
        "zip_entry_count": binding["archive"]["zip_entry_count"],
        "total_uncompressed_bytes": binding["archive"]["total_uncompressed_bytes"],
        "plan_sha256": binding["emulation_plan"]["sha256"],
        "plan_sheets": binding["emulation_plan"]["sheet_names"],
        "technique_ids_found": binding["emulation_plan"]["technique_ids_found"],
        "detector_executed": False,
    }, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
